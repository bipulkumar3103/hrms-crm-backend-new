import os
import subprocess
import sys

def install(package):
    subprocess.check_call([sys.executable, "-m", "pip", "install", package])

try:
    import docx
except ImportError:
    install('python-docx')
    import docx

from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

doc = Document()

# Add Title
title = doc.add_heading('Enterprise HRMS SaaS - End User Manual', 0)
title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

doc.add_paragraph("Note from Developer: You can drag and drop your screenshots directly below the camera placeholders in this document.\n")

# Section 1
doc.add_heading('1. Getting Started & Logging In', level=1)
doc.add_paragraph("Welcome to the Enterprise HRMS & CRM Platform. Our system is designed using 'Enterprise Elite' standards to ensure an intuitive and premium user experience.")
doc.add_heading('Accessing the Platform', level=2)
p1 = doc.add_paragraph("1. Navigate to the main login portal.\n2. Enter your authorized administrative credentials.\n3. If your account is tied to multiple organizational branches, you will be automatically routed based on your highest permission matrix.")
doc.add_paragraph("\n[ 📸 INSTRUCTION: Drop your 'Login Page' screenshot below this line ]", style='Intense Quote')
doc.add_paragraph("*(Showcase the refined accessibility and aesthetic entry pipeline)*")

# Section 2
doc.add_heading('2. Navigating the Elite Dashboard', level=1)
doc.add_paragraph("Upon logging in, you will be greeted by the primary dashboard. This interface is composed of autonomous 'Elite Blocks.'")
doc.add_heading('Dashboard Features', level=2)
doc.add_paragraph("- Dynamic Metrics: Each card on your dashboard independently fetches real-time data.", style='List Bullet')
doc.add_paragraph("- Customization Engine: Depending on your privileges, you can adjust typography, spacing, and thematic colors.", style='List Bullet')
doc.add_paragraph("\n[ 📸 INSTRUCTION: Drop your 'Main Dashboard Overview' screenshot below this line ]", style='Intense Quote')
doc.add_paragraph("*(Showcase the customizable UI blocks and the general layout)*")

# Section 3
doc.add_heading('3. The Live Network Organization Tree', level=1)
doc.add_paragraph("Our state-of-the-art Organizational Matrix allows administrators to visually manage company structures, regardless of size or depth.")
doc.add_heading('How to Use the Matrix', level=2)
doc.add_paragraph("1. Expanding Branches: Click on any node to expand it. The matrix features Intelligent Viewport Auto-Centering.")
doc.add_paragraph("2. Data Discovery Tooltips: Hover over any employee or department node to view rich, dynamic data.")
doc.add_paragraph("3. Zoom & Pan: Click and drag the canvas to explore different remote branches effortlessly.")
doc.add_paragraph("\n[ 📸 INSTRUCTION: Drop your 'Organization Tree' screenshot below this line ]", style='Intense Quote')
doc.add_paragraph("*(Showcase the tree, highlighting an active Data Discovery Tooltip on hover)*")

# Section 4
doc.add_heading('4. Project Governance & Timesheet Controls', level=1)
doc.add_heading('Managing the Project Registry', level=2)
doc.add_paragraph("- Interactive Accordions: Expand sectors to view hundreds of concurrent projects neatly.", style='List Bullet')
doc.add_paragraph("- Decommissioned Projects: Expired projects are highlighted and strictly prevented from being selected in new active timesheets.", style='List Bullet')
doc.add_paragraph("\n[ 📸 INSTRUCTION: Drop your 'Project Registry' screenshot below this line ]", style='Intense Quote')

doc.add_heading('Submitting Timesheets', level=2)
doc.add_paragraph("1. Go to the Timesheets module.\n2. Select your active project from the dropdown.\n3. Enter your hours and submit for matrix manager approval.")
doc.add_paragraph("\n[ 📸 INSTRUCTION: Drop your 'Timesheet Submission' screenshot below this line ]", style='Intense Quote')

# Section 5
doc.add_heading('5. Employee Profile Management & Administration', level=1)
doc.add_paragraph("Administrators have comprehensive control over user access and profile matrices.")
doc.add_paragraph("- Navigate to the Employee Admin section.\n- Apply overrides, shift team assignments, and manage system privileges seamlessly.")
doc.add_paragraph("\n[ 📸 INSTRUCTION: Drop your 'Employee Profile Management' screenshot below this line ]", style='Intense Quote')

doc.add_paragraph("\nEnd of Document.")

output_path = r'd:\Bipul-hrms\hrms-crm-backend-new\SaaS_User_Manual.docx'
doc.save(output_path)
print(f"Document generated successfully at {output_path}")
