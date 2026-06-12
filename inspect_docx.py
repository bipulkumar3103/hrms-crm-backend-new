from docx import Document

def inspect_docx(file_path):
    doc = Document(file_path)
    print(f"Total Paragraphs: {len(doc.paragraphs)}")
    print("-" * 30)
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            print(f"[{i}] {p.style.name}: {p.text[:100]}...")
        # Check for images in this paragraph
        for run in p.runs:
            if 'picture' in run.element.xml:
                print(f"  -> [IMAGE FOUND IN PARAGRAPH {i}]")

if __name__ == "__main__":
    inspect_docx("NexusOS_Enterprise_Technical_Manual.docx")
